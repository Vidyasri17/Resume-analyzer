import io
import re

def extract_text(filename: str, content: bytes) -> str:
    ext = filename.lower().rsplit(".",1)[-1] if "." in filename else ""
    if ext=="pdf":
        return extract_pdf(content)
    elif ext=="docx":
        return extract_docx(content)
    elif ext in ("txt","text","md"):
        return content.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file type: {ext}. Supported: PDF, DOCX, TXT")

def extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    texts=[]
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
            texts.append(t)
        except: continue
    text="\n".join(texts)
    if not text.strip():
        raise ValueError("PDF appears empty or scanned image (no extractable text)")
    return text

def extract_docx(content: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(content))
    parts=[]
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    text="\n".join(parts)
    if not text.strip():
        raise ValueError("DOCX appears empty")
    return text

def clean_text(text:str)->str:
    text=re.sub(r'\r\n','\n',text)
    text=re.sub(r'\n{3,}','\n\n',text)
    return text.strip()
